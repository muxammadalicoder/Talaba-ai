from docx import Document
from pptx import Presentation


def analyze(path: str) -> str:
    """Shablon tuzilmasini matn ko'rinishida qaytaradi (AI ga beriladi)."""
    if path.lower().endswith(".pptx"):
        p = Presentation(path)
        rows = []
        for i, s in enumerate(p.slides, 1):
            title = s.shapes.title.text if s.shapes.title is not None else ""
            body = [sh.text_frame.text[:60] for sh in s.shapes if sh.has_text_frame and sh != s.shapes.title and sh.text_frame.text.strip()]
            rows.append(f"{i}) [{s.slide_layout.name}] {title} | {' / '.join(body[:2])}")
        layouts = ", ".join(l.name for l in p.slide_layouts)
        return f"PPTX: {len(rows)} slayd. Maketlar: {layouts}. Slaydlar:\n" + "\n".join(rows[:30])
    d = Document(path)
    n = d.styles["Normal"]
    head = [f"{p.style.name}: {p.text.strip()[:80]}" for p in d.paragraphs
            if p.style.name.lower().startswith(("heading", "title", "sarlavha")) and p.text.strip()]
    plain = [p.text.strip()[:100] for p in d.paragraphs if p.text.strip() and p.style.name == "Normal"][:8]
    return (f"DOCX: shrift {n.font.name}, o'lcham {n.font.size.pt if n.font.size else '?'}, "
            f"{len(d.paragraphs)} abzas, {len(d.tables)} jadval.\nSarlavhalar tuzilmasi:\n" + "\n".join(head[:40]) +
            "\nMatn namunasi:\n" + "\n".join(plain))


import re

_NUM2 = re.compile(r"^\d+\.\d+")
_TOC = re.compile(r"(\.{3,}|\t\d+\s*$|\s\d{1,3}\s*$)")
_SPECIAL = re.compile(r"^\s*(kirish|xulosa|введение|заключение|introduction|conclusion)|adabiyot|литератур|references|ilova", re.I)
_START = re.compile(r"^\s*(kirish|введение|introduction)", re.I)


def is_special(title: str) -> bool:
    return bool(_SPECIAL.search(title))


def structure(path: str) -> dict:
    """Namuna kurs ishining sarlavhalari, hajmi (so'z) va turini aniqlaydi."""
    d = Document(path)
    paras = [p for p in d.paragraphs if p.text.strip()]

    def is_head(p):
        t = p.text.strip(); sn = p.style.name.lower()
        if len(t) >= 120 or _TOC.search(t):
            return False
        bold = bool(p.runs) and all(r.bold for r in p.runs if r.text.strip())
        numbered = _NUM2.match(t) or re.match(r"^\d+[-\s]*(bob|глава|chapter)", t, re.I)
        return sn.startswith(("heading", "title", "sarlavha")) or (
            not t.endswith((".", ",", ";")) and (bold or t.isupper() or numbered))

    start = next((i for i, p in enumerate(paras) if is_head(p) and _START.match(p.text)), 0)
    items = []
    for p in paras[start:]:
        t = p.text.strip()
        if is_head(p):
            lvl = 2 if (_NUM2.match(t) or p.style.name.lower().endswith("2")) else 1
            items.append({"level": lvl, "title": t, "words": 0})
        elif items:
            items[-1]["words"] += len(t.split())
    leaf = [it for k, it in enumerate(items)
            if not is_special(it["title"]) and (k + 1 == len(items) or items[k + 1]["level"] <= it["level"])]
    wps = int(sum(i["words"] for i in leaf) / len(leaf)) if leaf else 300
    head = " ".join(p.text for p in paras[:25]).lower()
    kind = "mustaqil" if "mustaqil" in head else "loyiha" if "loyiha" in head else "kurs"
    return {"kind": kind, "items": items, "total_words": sum(i["words"] for i in items),
            "wps": min(900, max(150, wps))}
