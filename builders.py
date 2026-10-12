import time
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt as PPt

NAVY, ACC, LIGHT = PRGB(0x14, 0x2B, 0x4F), PRGB(0x1F, 0x9E, 0x89), PRGB(0xF5, 0xF8, 0xFC)


# ---------------- PPTX ----------------
def _clear(prs):
    ids = prs.slides._sldIdLst
    for s in list(ids):
        prs.part.drop_rel(s.rId)
        ids.remove(s)


def _bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def _rect(slide, x, y, w, h, color):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background()


def _text(slide, x, y, w, h, lines, size, color, bold=False, gap=0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln; p.font.size = PPt(size); p.font.bold = bold; p.font.color.rgb = color
        p.space_after = PPt(gap)


def _title(slide, text, x, y, w, h, size, color):
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = x, y, w, h
    t.text = text
    t.text_frame.word_wrap = True
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT; p.font.size = PPt(size); p.font.bold = True; p.font.color.rgb = color


def build_pptx(plan, out, tpl=None):
    if tpl:
        prs = Presentation(tpl); _clear(prs)
        L = prs.slide_layouts
        s = prs.slides.add_slide(L[0]); s.shapes.title.text = plan["title"]
        if len(s.placeholders) > 1:
            s.placeholders[1].text = plan.get("subtitle", "")
        lay = L[1] if len(L) > 1 else L[0]
        for sl in plan["slides"]:
            s = prs.slides.add_slide(lay); s.shapes.title.text = sl["title"]
            body = next((p for p in s.placeholders if p.placeholder_format.idx == 1), None)
            if body is not None:
                tf = body.text_frame; tf.clear()
                for i, b in enumerate(sl["bullets"]):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = b
            if sl.get("notes"):
                s.notes_slide.notes_text_frame.text = sl["notes"]
        prs.save(out); return
    from deck import build_default
    build_default(plan, out)


# ---------------- DOCX ----------------
def _field(par, code):
    r = par.add_run()
    a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = code
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "end")
    r._r.extend([a, i, b])


def _style(doc):
    s = doc.sections[0]
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
    n = doc.styles["Normal"]
    n.font.name = "Times New Roman"; n.font.size = Pt(14)
    n.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    n.paragraph_format.line_spacing = 1.5; n.paragraph_format.space_after = Pt(0)
    for name, size in (("Heading 1", 16), ("Heading 2", 14)):
        h = doc.styles[name]
        h.font.name = "Times New Roman"; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = RGBColor(0, 0, 0)
        rf = h.element.rPr.rFonts
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(a), None)
        rf.set(qn("w:ascii"), "Times New Roman"); rf.set(qn("w:hAnsi"), "Times New Roman")
        h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(8)
        h.paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].paragraph_format.alignment = AL.CENTER


def _title_page(doc, plan, kind, prof):
    names = {"kurs": "KURS ISHI", "mustaqil": "MUSTAQIL ISH", "loyiha": "KURS LOYIHASI"}

    def c(text, size=14, bold=False, after=0, align=AL.CENTER):
        p = doc.add_paragraph(); p.alignment = align
        p.paragraph_format.first_line_indent = 0; p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)

    c("O'ZBEKISTON RESPUBLIKASI OLIY TA'LIM, FAN VA INNOVATSIYALAR VAZIRLIGI", bold=True, after=12)
    c(prof.get("university", "[Universitet nomi]"), bold=True)
    c(prof.get("faculty", "[Fakultet]"), after=90)
    c(names.get(kind, "ISH"), 20, True, 24)
    c("Mavzu: " + plan["title"], 16, True, 90)
    c("Bajardi: " + prof.get("name", "[F.I.Sh.]"), align=AL.RIGHT)
    c("Tekshirdi: " + prof.get("teacher", "[O'qituvchi]"), after=90, align=AL.RIGHT)
    c(f"{prof.get('city', 'Toshkent')} – {time.strftime('%Y')}")


def build_docx(plan, blocks, out, tpl=None, kind="kurs", prof=None):
    prof = prof or {}
    if tpl:
        doc = Document(tpl)
        body = doc.element.body
        for el in list(body):
            if not el.tag.endswith("sectPr"):
                body.remove(el)
        nf = doc.styles["Normal"].font.name or "Times New Roman"
        for name in ("Heading 1", "Heading 2"):
            try:
                h = doc.styles[name]
            except KeyError:
                continue
            h.font.name = nf; h.font.bold = True; h.font.color.rgb = RGBColor(0, 0, 0)
            rf = h.element.rPr.rFonts
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rf.attrib.pop(qn(a), None)
            rf.set(qn("w:ascii"), nf); rf.set(qn("w:hAnsi"), nf)
    else:
        doc = Document(); _style(doc)
    if kind != "task":
        _title_page(doc, plan, kind, prof); doc.add_page_break()
        p = doc.add_paragraph(); p.alignment = AL.CENTER
        p.add_run("MUNDARIJA").bold = True
        _field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
        doc.add_page_break()
        st = OxmlElement("w:updateFields"); st.set(qn("w:val"), "true"); doc.settings.element.append(st)
        f = doc.sections[0].footer.paragraphs[0]; f.alignment = AL.CENTER; _field(f, "PAGE")
    refs = 0
    for b in blocks:
        t, x = b["t"], b["x"]
        if t in ("h1", "h2"):
            doc.add_heading(x, 1 if t == "h1" else 2)
        elif t == "ref":
            refs += 1
            p = doc.add_paragraph(f"{refs}. {x}"); p.alignment = AL.JUSTIFY
        else:
            for ln in x.split("\n"):
                if not ln.strip():
                    continue
                if ln.startswith("- "):
                    try:
                        doc.add_paragraph(ln[2:], style="List Bullet")
                        continue
                    except KeyError:
                        pass
                p = doc.add_paragraph(ln); p.alignment = AL.JUSTIFY
                p.paragraph_format.first_line_indent = Cm(1.25)
    doc.save(out)
