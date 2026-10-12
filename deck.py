import io
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor as C
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as I, Pt
import builders as B

THEMES = {"navy": ("142B4F", "1F9E89", "F5F8FC"), "emerald": ("0F3D3A", "F2A93B", "F4FAF8"),
          "crimson": ("3B1620", "E4572E", "FBF6F5"), "violet": ("2A1B54", "7C5CFF", "F7F5FF")}
SOFT = C.from_string("C9D6EA"); GRAY = C.from_string("6B7280"); WHITE = C(255, 255, 255)


def _pic(s, data, x, y, w, h):
    iw, ih = Image.open(io.BytesIO(data)).size
    pic = s.shapes.add_picture(io.BytesIO(data), x, y, w, h)
    br, ir = w / h, iw / ih
    if ir > br:
        c = (1 - br / ir) / 2; pic.crop_left = pic.crop_right = c
    else:
        c = (1 - ir / br) / 2; pic.crop_top = pic.crop_bottom = c


def _bullets(s, items, x, y, w, h, ink, size=None):
    tot = sum(len(t) for t in items)
    B._text(s, x, y, w, h, ["•  " + t for t in items], size or (24 if tot < 260 else 20 if tot < 420 else 17), ink, gap=12)


def _chart(s, ch, x, y, w, h, acc):
    cd = CategoryChartData(); cd.categories = [str(c) for c in ch["categories"]]
    for se in ch["series"]:
        vals = [float(v) for v in se["values"]]
        if len(vals) != len(ch["categories"]):
            raise ValueError("uzunlik mos emas")
        cd.add_series(se.get("name", ""), vals)
    col, line, pie = XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.PIE
    t = {"bar": col, "line": line, "pie": pie}.get(ch.get("type"), col)
    c = s.shapes.add_chart(t, x, y, w, h, cd).chart
    c.font.size = Pt(13); c.has_title = False
    c.has_legend = len(ch["series"]) > 1 or t == pie
    if c.has_legend:
        c.legend.position = XL_LEGEND_POSITION.BOTTOM; c.legend.include_in_layout = False
    if t != pie and len(ch["series"]) == 1:
        ser = c.plots[0].series[0]
        if t == col:
            ser.format.fill.solid(); ser.format.fill.fore_color.rgb = acc
        else:
            ser.format.line.color.rgb = acc


def _process(s, items, acc, ink):
    items = items[:5]; n = len(items); gap = 0.3; w = (11.6 - gap * (n - 1)) / n
    for k, t in enumerate(items):
        x = 0.9 + k * (w + gap)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, I(x), I(2.1), I(w), I(4.2))
        card.fill.solid(); card.fill.fore_color.rgb = WHITE; card.line.color.rgb = acc
        o = s.shapes.add_shape(MSO_SHAPE.OVAL, I(x + 0.25), I(2.35), I(0.7), I(0.7))
        o.fill.solid(); o.fill.fore_color.rgb = acc; o.line.fill.background()
        o.text_frame.text = str(k + 1); p = o.text_frame.paragraphs[0]
        p.font.size = Pt(22); p.font.bold = True; p.font.color.rgb = WHITE; p.alignment = PP_ALIGN.CENTER
        B._text(s, I(x + 0.15), I(3.3), I(w - 0.3), I(2.9), [t], 16 if n > 3 else 18, ink)


def build_default(plan, out):
    dark, acc, light = [C.from_string(h) for h in THEMES.get(plan.get("theme"), THEMES["navy"])]
    ink = C.from_string("232B38")
    prs = Presentation(); prs.slide_width, prs.slide_height = I(13.333), I(7.5)
    lay = prs.slide_layouts[5]
    s = prs.slides.add_slide(lay); B._bg(s, dark)
    cov = plan.get("_cover"); tw = 11.5
    if cov:
        _pic(s, cov[0], I(7.6), 0, I(5.733), I(7.5)); B._rect(s, I(7.45), 0, I(0.15), I(7.5), acc); tw = 6.3
    B._title(s, plan["title"], I(0.9), I(2.2), I(tw), I(2.2), 38 if cov else 44, WHITE)
    B._rect(s, I(0.9), I(4.5), I(2), I(0.08), acc)
    B._text(s, I(0.9), I(4.75), I(tw), I(1.2), [plan.get("subtitle", "")], 20, SOFT)
    if cov:
        s.notes_slide.notes_text_frame.text = "Muqova rasmi: " + cov[1]

    def base(title, n):
        s = prs.slides.add_slide(lay); B._bg(s, light); B._rect(s, 0, 0, I(0.3), prs.slide_height, acc)
        B._title(s, title, I(0.8), I(0.35), I(11.8), I(1.1), 30 if len(title) < 60 else 24, dark)
        B._text(s, I(12.2), I(6.95), I(0.9), I(0.4), [str(n)], 12, acc)
        return s

    for n, sl in enumerate(plan["slides"], 1):
        s = base(sl["title"], n)
        b = sl.get("bullets", []); lt = sl.get("layout", "text"); img = sl.get("_img"); done = False
        if lt == "stat" and sl.get("stat"):
            st = sl["stat"]; B._rect(s, I(0.9), I(1.8), I(4.6), I(4.4), dark)
            B._text(s, I(1.1), I(2.4), I(4.2), I(1.6), [str(st.get("value", ""))], 60, WHITE, bold=True)
            B._text(s, I(1.1), I(4.2), I(4.2), I(1.8), [st.get("label", "")], 18, SOFT)
            _bullets(s, b, I(6.0), I(1.9), I(6.8), I(4.8), ink); done = True
        elif lt == "chart" and sl.get("chart"):
            try:
                _chart(s, sl["chart"], I(0.9), I(1.7), I(7.2), I(5.0), acc)
                _bullets(s, b[:3], I(8.4), I(1.8), I(4.4), I(4.8), ink, 18); done = True
            except Exception:
                pass
        elif lt == "process" and len(b) >= 3:
            _process(s, b, acc, ink); done = True
        elif lt == "two_col" and len(b) >= 4:
            h = (len(b) + 1) // 2
            _bullets(s, b[:h], I(0.9), I(1.8), I(5.7), I(4.9), ink, 20)
            B._rect(s, I(6.75), I(1.9), I(0.04), I(4.6), acc)
            _bullets(s, b[h:], I(7.0), I(1.8), I(5.7), I(4.9), ink, 20); done = True
        if not done:
            if img:
                _bullets(s, b, I(0.9), I(1.7), I(6.9), I(5.0), ink)
                B._rect(s, I(8.32), I(1.82), I(4.6), I(4.9), acc)
                _pic(s, img[0], I(8.2), I(1.7), I(4.6), I(4.9))
                B._text(s, I(8.2), I(6.75), I(4.6), I(0.3), ["Rasm: " + img[1]], 9, GRAY)
            else:
                _bullets(s, b, I(0.9), I(1.7), I(11.5), I(5.0), ink)
        notes = sl.get("notes", "")
        if img:
            notes += f"\n\n[Rasm manbasi: {img[1]}]"
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    src = [x for x in plan.get("sources", []) if x][:8]
    if src:
        s = base("Foydalanilgan manbalar", len(plan["slides"]) + 1)
        B._text(s, I(0.9), I(1.7), I(11.5), I(5), src, 14, ink, gap=8)
    prs.save(out)
